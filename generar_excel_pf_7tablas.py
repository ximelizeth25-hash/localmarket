"""
Script Profesional de Métricas de Software: Análisis de Puntos de Función (IFPUG / Albrecht)
Generación de Libro Excel Automatizado basado en plantilla 'PuntosDeFuncion.DOC'
Mapeo exhaustivo de las 7 tablas de la base de datos:
1. Alumno
2. Estado
3. Pago
4. Materia
5. Cursa
6. Becado
7. Discapacitado
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generar_excel_puntos_de_funcion():
    wb = openpyxl.Workbook()
    # Remover hoja por defecto
    default_sheet = wb.active

    # =========================================================================
    # PALETA DE ESTILOS PROFESIONALES (Azul #1F4E78, Totales Verdes, Cuadrícula)
    # =========================================================================
    COLOR_AZUL_HDR = "1F4E78"       # Azul corporativo solicitado
    COLOR_AZUL_SUB = "2F5597"       # Azul secundario para subsecciones
    COLOR_AZUL_LIGHT = "D9E1F2"     # Azul suave para subtotales
    COLOR_VERDE_TOT = "E2F0D9"      # Verde pastel para totales
    COLOR_VERDE_TEXT = "1E4620"     # Verde oscuro para texto de totales
    COLOR_VERDE_ACCENT = "385723"   # Verde fuerte para bordes de total
    COLOR_ZEBRA = "F9FAFB"          # Fondo alterno suave
    COLOR_BORDER = "D9D9D9"         # Borde gris sutil

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="404040")
    font_section = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_tbl_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="000000")
    font_data_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_subtotal = Font(name="Calibri", size=10, bold=True, color="1F4E78")
    font_total_green = Font(name="Calibri", size=11, bold=True, color=COLOR_VERDE_TEXT)
    font_kpi_val = Font(name="Calibri", size=18, bold=True, color="1F4E78")
    font_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="595959")

    fill_title = PatternFill(start_color=COLOR_AZUL_HDR, end_color=COLOR_AZUL_HDR, fill_type="solid")
    fill_section = PatternFill(start_color=COLOR_AZUL_SUB, end_color=COLOR_AZUL_SUB, fill_type="solid")
    fill_tbl_header = PatternFill(start_color=COLOR_AZUL_HDR, end_color=COLOR_AZUL_HDR, fill_type="solid")
    fill_subtotal = PatternFill(start_color=COLOR_AZUL_LIGHT, end_color=COLOR_AZUL_LIGHT, fill_type="solid")
    fill_total_green = PatternFill(start_color=COLOR_VERDE_TOT, end_color=COLOR_VERDE_TOT, fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
    fill_kpi = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    thin_border = Side(style="thin", color=COLOR_BORDER)
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    
    border_subtotal = Border(
        top=Side(style="thin", color="1F4E78"),
        bottom=Side(style="double", color="1F4E78"),
        left=thin_border,
        right=thin_border
    )

    border_total_green = Border(
        top=Side(style="thin", color=COLOR_VERDE_ACCENT),
        bottom=Side(style="double", color=COLOR_VERDE_ACCENT),
        left=thin_border,
        right=thin_border
    )

    # =========================================================================
    # HOJA 1: MATRICES DE COMPLEJIDAD (Plantilla Oficial del Documento Word)
    # =========================================================================
    ws_mat = wb.create_sheet(title="Matrices de Complejidad")
    ws_mat.views.sheetView[0].showGridLines = True

    # Encabezado
    ws_mat.merge_cells("A1:G1")
    t1 = ws_mat["A1"]
    t1.value = "MATRICES OFICIALES DE COMPLEJIDAD Y VALORACIÓN (IFPUG / ALBRECHT)"
    t1.font = font_title
    t1.fill = fill_title
    t1.alignment = Alignment(horizontal="center", vertical="center")
    ws_mat.row_dimensions[1].height = 36

    ws_mat.merge_cells("A2:G2")
    s1 = ws_mat["A2"]
    s1.value = "Reglas de ponderación y tablas de referencia estandarizadas de 'PuntosDeFuncion.DOC'"
    s1.font = font_subtitle
    s1.alignment = Alignment(horizontal="center", vertical="center")
    ws_mat.row_dimensions[2].height = 20

    row_m = 4

    # Definiciones Fundamentales
    ws_mat.merge_cells(start_row=row_m, start_column=1, end_row=row_m, end_column=7)
    sec_def = ws_mat.cell(row=row_m, column=1, value="1. DEFINICIONES OPERATIVAS DEL MODELO")
    sec_def.font = font_section
    sec_def.fill = fill_section
    sec_def.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_mat.row_dimensions[row_m].height = 24
    row_m += 1

    defs_data = [
        ("Entrada (EI)", "Transacciones elementales relacionadas con agregar (altas), eliminar (bajas físicas o lógicas) y modificar datos en las tablas."),
        ("Salida (EO)", "Consultas o generación de informes que involucran cálculos derivados, operaciones matemáticas, prorrateos o saldos acumulados."),
        ("Petición (EQ)", "Consultas directas o generación de reportes sin cálculos matemáticos derivados (como listados de catálogos y búsquedas)."),
        ("Archivos Internos (ILF)", "Archivos lógicos o tablas que son mantenidos y alterados directamente por el propio sistema (Alumno, Pago, Materia, etc.)."),
        ("Archivos Externos (EIF)", "Archivos referenciados solo para consulta por el sistema pero alterados por sistemas externos (RENAPO, SEJ).")
    ]

    for item, desc in defs_data:
        ws_mat.cell(row=row_m, column=1, value=item).font = font_data_bold
        ws_mat.cell(row=row_m, column=1).border = border_cell
        ws_mat.cell(row=row_m, column=1).alignment = Alignment(horizontal="left", vertical="center")
        
        ws_mat.merge_cells(start_row=row_m, start_column=2, end_row=row_m, end_column=7)
        c_desc = ws_mat.cell(row=row_m, column=2, value=desc)
        c_desc.font = font_data
        c_desc.border = border_cell
        c_desc.alignment = Alignment(horizontal="left", vertical="center")
        for col in range(3, 8):
            ws_mat.cell(row=row_m, column=col).border = border_cell
        ws_mat.row_dimensions[row_m].height = 20
        row_m += 1

    row_m += 1

    # Tabla 1: Matriz para Entradas
    ws_mat.merge_cells(start_row=row_m, start_column=1, end_row=row_m, end_column=7)
    sec_t1 = ws_mat.cell(row=row_m, column=1, value="TABLA 1. MATRIZ DE COMPLEJIDAD PARA ENTRADAS EXTERNAS (EI)")
    sec_t1.font = font_section
    sec_t1.fill = fill_section
    sec_t1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_mat.row_dimensions[row_m].height = 24
    row_m += 1

    t1_headers = ["Archivos Referenciados (FTR)", "1 – 4 Elementos (DET)", "5 – 15 Elementos (DET)", "16 o más Elementos (DET)"]
    for c_i, h in enumerate(t1_headers, 1):
        cell_h = ws_mat.cell(row=row_m, column=c_i, value=h)
        cell_h.font = font_tbl_header
        cell_h.fill = fill_tbl_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = border_cell
    ws_mat.row_dimensions[row_m].height = 22
    row_m += 1

    t1_rows = [
        ["Menos de 2 (< 2 FTR)", "Bajo", "Bajo", "Medio"],
        ["2 FTR", "Bajo", "Medio", "Alto"],
        ["Más de 2 (> 2 FTR)", "Medio", "Alto", "Alto"]
    ]
    for r in t1_rows:
        for c_i, val in enumerate(r, 1):
            cell = ws_mat.cell(row=row_m, column=c_i, value=val)
            cell.font = font_data_bold if c_i == 1 else font_data
            cell.alignment = Alignment(horizontal="left" if c_i == 1 else "center", vertical="center")
            cell.border = border_cell
        ws_mat.row_dimensions[row_m].height = 19
        row_m += 1

    row_m += 1

    # Tabla 2: Matriz para Salidas y Peticiones
    ws_mat.merge_cells(start_row=row_m, start_column=1, end_row=row_m, end_column=7)
    sec_t2 = ws_mat.cell(row=row_m, column=1, value="TABLA 2. MATRIZ DE COMPLEJIDAD PARA SALIDAS (EO) Y PETICIONES (EQ)")
    sec_t2.font = font_section
    sec_t2.fill = fill_section
    sec_t2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_mat.row_dimensions[row_m].height = 24
    row_m += 1

    t2_headers = ["Archivos Referenciados (FTR)", "1 – 5 Elementos (DET)", "6 – 19 Elementos (DET)", "20 o más Elementos (DET)"]
    for c_i, h in enumerate(t2_headers, 1):
        cell_h = ws_mat.cell(row=row_m, column=c_i, value=h)
        cell_h.font = font_tbl_header
        cell_h.fill = fill_tbl_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = border_cell
    ws_mat.row_dimensions[row_m].height = 22
    row_m += 1

    t2_rows = [
        ["Menos que 2 (< 2 FTR)", "Bajo", "Bajo", "Medio"],
        ["2 – 3 FTR", "Bajo", "Medio", "Alto"],
        ["Más de 3 (> 3 FTR)", "Medio", "Alto", "Alto"]
    ]
    for r in t2_rows:
        for c_i, val in enumerate(r, 1):
            cell = ws_mat.cell(row=row_m, column=c_i, value=val)
            cell.font = font_data_bold if c_i == 1 else font_data
            cell.alignment = Alignment(horizontal="left" if c_i == 1 else "center", vertical="center")
            cell.border = border_cell
        ws_mat.row_dimensions[row_m].height = 19
        row_m += 1

    row_m += 1

    # Tabla 3: Matriz para Archivos Internos y Externos
    ws_mat.merge_cells(start_row=row_m, start_column=1, end_row=row_m, end_column=7)
    sec_t3 = ws_mat.cell(row=row_m, column=1, value="TABLA 3. MATRIZ DE COMPLEJIDAD PARA ARCHIVOS INTERNOS (ILF) Y EXTERNOS (EIF)")
    sec_t3.font = font_section
    sec_t3.fill = fill_section
    sec_t3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_mat.row_dimensions[row_m].height = 24
    row_m += 1

    t3_headers = ["Tipos de Registro (RET)", "1 – 19 Elementos (DET)", "20 – 50 Elementos (DET)", "51 o más Elementos (DET)"]
    for c_i, h in enumerate(t3_headers, 1):
        cell_h = ws_mat.cell(row=row_m, column=c_i, value=h)
        cell_h.font = font_tbl_header
        cell_h.fill = fill_tbl_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = border_cell
    ws_mat.row_dimensions[row_m].height = 22
    row_m += 1

    t3_rows = [
        ["Menos de 2 (1 RET)", "Bajo", "Bajo", "Medio"],
        ["2 – 5 RET", "Bajo", "Medio", "Alto"],
        ["Más de 5 (> 5 RET)", "Medio", "Alto", "Alto"]
    ]
    for r in t3_rows:
        for c_i, val in enumerate(r, 1):
            cell = ws_mat.cell(row=row_m, column=c_i, value=val)
            cell.font = font_data_bold if c_i == 1 else font_data
            cell.alignment = Alignment(horizontal="left" if c_i == 1 else "center", vertical="center")
            cell.border = border_cell
        ws_mat.row_dimensions[row_m].height = 19
        row_m += 1

    row_m += 1

    # Tabla 4: Puntos de Función No Ajustados (Pesos)
    ws_mat.merge_cells(start_row=row_m, start_column=1, end_row=row_m, end_column=7)
    sec_t4 = ws_mat.cell(row=row_m, column=1, value="TABLA 4. TABLA DE PONDERACIÓN OFICIAL DE PUNTOS DE FUNCIÓN NO AJUSTADOS")
    sec_t4.font = font_section
    sec_t4.fill = fill_section
    sec_t4.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_mat.row_dimensions[row_m].height = 24
    row_m += 1

    t4_headers = ["Componente Funcional", "Bajo", "Medio", "Alto"]
    for c_i, h in enumerate(t4_headers, 1):
        cell_h = ws_mat.cell(row=row_m, column=c_i, value=h)
        cell_h.font = font_tbl_header
        cell_h.fill = fill_tbl_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center")
        cell_h.border = border_cell
    ws_mat.row_dimensions[row_m].height = 22
    row_m += 1

    t4_rows = [
        ["Entradas Externas (EI)", 3, 4, 6],
        ["Salidas Externas (EO)", 4, 5, 7],
        ["Peticiones / Consultas (EQ)", 3, 4, 6],
        ["Archivos Lógicos Internos (ILF)", 7, 10, 15],
        ["Archivos de Interfaz Externa (EIF)", 5, 7, 10]
    ]
    for r in t4_rows:
        for c_i, val in enumerate(r, 1):
            cell = ws_mat.cell(row=row_m, column=c_i, value=val)
            cell.font = font_data_bold if c_i == 1 else font_data
            cell.alignment = Alignment(horizontal="left" if c_i == 1 else "center", vertical="center")
            cell.border = border_cell
        ws_mat.row_dimensions[row_m].height = 19
        row_m += 1

    # Autoajuste de anchos para Hoja 1
    col_w_m = {1: 36, 2: 26, 3: 26, 4: 26, 5: 16, 6: 16, 7: 16}
    for c_idx, width in col_w_m.items():
        ws_mat.column_dimensions[get_column_letter(c_idx)].width = width

    # =========================================================================
    # HOJA 2: PUNTOS DE FUNCIÓN NO AJUSTADOS (PFNA) - DESGLOSE COMPLETO
    # =========================================================================
    ws_pfna = wb.create_sheet(title="PF No Ajustados")
    ws_pfna.views.sheetView[0].showGridLines = True

    # Encabezado Principal
    ws_pfna.merge_cells("A1:G1")
    t2 = ws_pfna["A1"]
    t2.value = "MATRIZ DE PUNTOS DE FUNCIÓN NO AJUSTADOS (PFNA) - 7 TABLAS BD"
    t2.font = font_title
    t2.fill = fill_title
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws_pfna.row_dimensions[1].height = 36

    ws_pfna.merge_cells("A2:G2")
    s2 = ws_pfna["A2"]
    s2.value = "Desglose de Altas, Bajas, Modificaciones, Reportes y Archivos para: Alumno, Estado, Pago, Materia, Cursa, Becado, Discapacitado"
    s2.font = font_subtitle
    s2.alignment = Alignment(horizontal="center", vertical="center")
    ws_pfna.row_dimensions[2].height = 20

    row_p = 4
    subtotal_rows_pfna = []

    def agregar_bloque_pfna(titulo, cabeceras, datos):
        nonlocal row_p
        # Header de Sección
        ws_pfna.merge_cells(start_row=row_p, start_column=1, end_row=row_p, end_column=7)
        sec = ws_pfna.cell(row=row_p, column=1, value=titulo)
        sec.font = font_section
        sec.fill = fill_section
        sec.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_pfna.row_dimensions[row_p].height = 24
        row_p += 1

        # Cabeceras
        for c_idx, h in enumerate(cabeceras, 1):
            c = ws_pfna.cell(row=row_p, column=c_idx, value=h)
            c.font = font_tbl_header
            c.fill = fill_tbl_header
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border_cell
        ws_pfna.row_dimensions[row_p].height = 24
        row_p += 1

        start_r = row_p
        for idx, r_data in enumerate(datos):
            is_even = (idx % 2 == 1)
            for c_idx, val in enumerate(r_data, 1):
                c = ws_pfna.cell(row=row_p, column=c_idx, value=val)
                c.font = font_data
                c.border = border_cell
                if is_even:
                    c.fill = fill_zebra
                
                if c_idx in [1, 2]:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="center", vertical="center")
            ws_pfna.row_dimensions[row_p].height = 20
            row_p += 1
        end_r = row_p - 1

        # Subtotal dinámico con fórmula SUM
        ws_pfna.merge_cells(start_row=row_p, start_column=1, end_row=row_p, end_column=6)
        sub_lbl = ws_pfna.cell(row=row_p, column=1, value=f"SUBTOTAL {titulo.split(' (')[0]}:")
        sub_lbl.font = font_subtotal
        sub_lbl.fill = fill_subtotal
        sub_lbl.alignment = Alignment(horizontal="right", vertical="center")
        sub_lbl.border = border_subtotal

        for col in range(2, 7):
            ws_pfna.cell(row=row_p, column=col).border = border_subtotal
            ws_pfna.cell(row=row_p, column=col).fill = fill_subtotal

        sub_val = ws_pfna.cell(row=row_p, column=7)
        sub_val.value = f"=SUM(G{start_r}:G{end_r})"
        sub_val.font = font_subtotal
        sub_val.fill = fill_subtotal
        sub_val.alignment = Alignment(horizontal="center", vertical="center")
        sub_val.border = border_subtotal
        ws_pfna.row_dimensions[row_p].height = 22

        subtotal_rows_pfna.append(row_p)
        row_p += 2

    # 1. ENTRADAS EXTERNAS (EI) - Altas, Bajas y Modificaciones
    datos_ei = [
        # Tabla Alumno
        ["Alta Alumno", "Registro de nuevo alumno en tabla Alumno con clave de estado", "2 (Alumno, Estado)", "7 (Matrícula, Clave_Estado, Nombre...)", "Medio", 1, 4],
        ["Modificar Alumno", "Actualización de datos personales y estado de procedencia", "2 (Alumno, Estado)", "7 (Nombre, Apellidos, Sexo, Fecha_Nac...)", "Medio", 1, 4],
        ["Baja Alumno", "Eliminación o baja lógica del estudiante", "1 (Alumno)", "1 (Matrícula)", "Bajo", 1, 3],
        # Tabla Estado
        ["Alta Estado", "Registro de nueva entidad federativa en catálogo Estado", "1 (Estado)", "2 (Clave_Estado, Nombre_Estado)", "Bajo", 1, 3],
        ["Modificar Estado", "Actualización del nombre de la entidad federativa", "1 (Estado)", "2 (Clave_Estado, Nombre_Estado)", "Bajo", 1, 3],
        ["Baja Estado", "Eliminación de estado no referenciado", "1 (Estado)", "1 (Clave_Estado)", "Bajo", 1, 3],
        # Tabla Pago
        ["Registrar Pago", "Captura de transacción de pago asociada al alumno", "2 (Pago, Alumno)", "4 (Número_Pago, Fecha, Concepto, Monto)", "Bajo", 1, 3],
        ["Modificar Pago", "Ajuste o corrección de fecha, concepto o monto de pago", "2 (Pago, Alumno)", "4 (Número_Pago, Fecha, Concepto, Monto)", "Bajo", 1, 3],
        ["Cancelar Pago", "Anulación y baja de comprobante de pago", "1 (Pago)", "1 (Número_Pago)", "Bajo", 1, 3],
        # Tabla Materia
        ["Alta Materia", "Registro de asignatura en catálogo Materia", "1 (Materia)", "3 (Clave_Materia, Nombre_Materia, Costo)", "Bajo", 1, 3],
        ["Modificar Materia", "Actualización de nombre y arancel/costo de la materia", "1 (Materia)", "3 (Clave_Materia, Nombre_Materia, Costo)", "Bajo", 1, 3],
        ["Baja Materia", "Eliminación de materia del plan de estudios", "1 (Materia)", "1 (Clave_Materia)", "Bajo", 1, 3],
        # Tabla Cursa
        ["Inscribir Cursa", "Alta de inscripción académica alumno-materia en Cursa", "2 (Cursa, Materia)", "3 (Matrícula, Clave_Materia, Fecha)", "Bajo", 1, 3],
        ["Modificar Cursa", "Cambio de materia o fecha de inscripción académica", "2 (Cursa, Materia)", "3 (Matrícula, Clave_Materia, Fecha)", "Bajo", 1, 3],
        ["Baja Cursa", "Retiro o desincorporación de materia cursada", "1 (Cursa)", "2 (Matrícula, Clave_Materia)", "Bajo", 1, 3],
        # Tabla Becado
        ["Asignar Beca", "Registro de beneficio económico y porcentaje en Becado", "2 (Becado, Alumno)", "3 (Matrícula, Porcentaje, Motivo)", "Bajo", 1, 3],
        ["Modificar Beca", "Actualización de porcentaje de descuento o motivo", "2 (Becado, Alumno)", "3 (Matrícula, Porcentaje, Motivo)", "Bajo", 1, 3],
        ["Eliminar Beca", "Revocación o conclusión de beneficio de beca", "1 (Becado)", "1 (Matrícula)", "Bajo", 1, 3],
        # Tabla Discapacitado
        ["Registrar Discapacitado", "Alta de condición especial y aparato ortopédico requerido", "2 (Discapacitado, Alumno)", "3 (Matrícula, Descripción, Aparato)", "Bajo", 1, 3],
        ["Modificar Discapacitado", "Actualización del diagnóstico o aparato de asistencia", "2 (Discapacitado, Alumno)", "3 (Matrícula, Descripción, Aparato)", "Bajo", 1, 3],
        ["Eliminar Discapacitado", "Baja del registro de requerimiento ortopédico", "1 (Discapacitado)", "1 (Matrícula)", "Bajo", 1, 3],
    ]
    agregar_bloque_pfna("1. ENTRADAS EXTERNAS (EI) - Altas, Bajas y Modificaciones por Tabla",
                       ["Transacción / Entrada", "Descripción del Proceso Elemental", "Archivos Ref. (FTR)", "Campos (DET)", "Complejidad", "Total Ops", "Puntos de Función"],
                       datos_ei)

    # 2. SALIDAS EXTERNAS (EO) - Con Cálculos Matemáticos Derivados
    datos_eo = [
        ["Reporte de Estado de Cuenta y Saldos", "Cálculo de suma de materias inscritas vs pagos efectuados y saldo deudor", "4 (Alumno, Pago, Cursa, Becado)", "8 (Matrícula, Nombre, Total_Cargos, Total_Abonos, Saldo...)", "Alto", 1, 7],
        ["Carga Académica con Costos", "Liquidación arancelaria de materias cursadas con deducción porcentual de beca", "3 (Cursa, Materia, Becado)", "7 (Matrícula, Materias, Costo_Base, Beca_Desc, Total_Pagar...)", "Medio", 1, 5],
        ["Informe Estadístico de Becas", "Cálculo del monto total subsidiado por la institución y porcentaje promedio", "3 (Becado, Pago, Alumno)", "6 (Motivo, Alumnos_Beneficiados, Descuento_Promedio, Subsidio...)", "Medio", 1, 5],
        ["Reporte Financiero de Ingresos", "Consolidado de cobros realizados por periodo, cálculo de IVA y balance general", "3 (Pago, Cursa, Materia)", "6 (Periodo, Total_Recaudado, Pagos_Emitidos, Promedio_Pago...)", "Medio", 1, 5]
    ]
    agregar_bloque_pfna("2. SALIDAS EXTERNAS (EO) - Con Cálculos Matemáticos Derivados",
                       ["Salida / Reporte Derivado", "Descripción del Cálculo y Proceso", "Archivos Ref. (FTR)", "Campos (DET)", "Complejidad", "Total Ops", "Puntos de Función"],
                       datos_eo)

    # 3. PETICIONES / CONSULTAS (EQ) - Sin Cálculos Derivados
    datos_eq = [
        ["Consulta de Alumno por Matrícula", "Búsqueda directa del expediente del alumno y nombre del estado de procedencia", "2 (Alumno, Estado)", "7 (Matrícula, Nombre, Apellidos, Sexo, Estado...)", "Medio", 1, 4],
        ["Catálogo de Materias", "Listado directo de materias activas y sus aranceles asignados", "1 (Materia)", "3 (Clave_Materia, Nombre_Materia, Costo)", "Bajo", 1, 3],
        ["Historial de Pagos de Alumno", "Listado cronológico de recibos emitidos por estudiante", "2 (Pago, Alumno)", "4 (Número_Pago, Fecha, Concepto, Monto)", "Bajo", 1, 3],
        ["Padrón de Alumnos Becados", "Consulta directa de estudiantes con beca y motivo asignado", "2 (Becado, Alumno)", "4 (Matrícula, Nombre_Alumno, Porcentaje, Motivo)", "Bajo", 1, 3],
        ["Padrón de Alumnos Discapacitados", "Consulta directa de estudiantes con necesidades y aparatos ortopédicos", "2 (Discapacitado, Alumno)", "4 (Matrícula, Nombre_Alumno, Discapacidad, Aparato)", "Bajo", 1, 3],
        ["Catálogo de Estados", "Consulta directa de estados de la república registrados", "1 (Estado)", "2 (Clave_Estado, Nombre_Estado)", "Bajo", 1, 3],
        ["Consulta de Cursa por Estudiante", "Listado de materias en las que está inscrito el estudiante", "2 (Cursa, Materia)", "4 (Matrícula, Clave_Materia, Nombre_Materia, Fecha)", "Bajo", 1, 3],
    ]
    agregar_bloque_pfna("3. PETICIONES / CONSULTAS EXTERNAS (EQ) - Sin Cálculos Derivados",
                       ["Petición / Consulta Directa", "Descripción de la Consulta", "Archivos Ref. (FTR)", "Campos (DET)", "Complejidad", "Total Ops", "Puntos de Función"],
                       datos_eq)

    # 4. ARCHIVOS LÓGICOS INTERNOS (ILF) - Mapeo de las 7 Tablas
    datos_ilf = [
        ["1. Tabla: Alumno", "Almacena los datos maestros y de identidad del estudiante", "1 (Registro Alumno)", "7 (Matrícula, Clave_Estado, Nombre, Ap_Pat, Ap_Mat, Sexo, Fecha_Nac)", "Bajo", 1, 7],
        ["2. Tabla: Estado", "Catálogo maestro de entidades federativas", "1 (Registro Estado)", "2 (Clave_Estado, Nombre_Estado)", "Bajo", 1, 7],
        ["3. Tabla: Pago", "Registro transaccional de pagos escolares efectuados", "1 (Registro Pago)", "4 (Número_Pago, Fecha, Concepto, Monto)", "Bajo", 1, 7],
        ["4. Tabla: Materia", "Catálogo de asignaturas curriculares y costos", "1 (Registro Materia)", "3 (Clave_Materia, Nombre_Materia, Costo)", "Bajo", 1, 7],
        ["5. Tabla: Cursa", "Entidad de relación muchos a muchos de inscripciones", "1 (Registro Cursa)", "3 (Matrícula, Clave_Materia, Fecha)", "Bajo", 1, 7],
        ["6. Tabla: Becado", "Entidad de soporte de asignación de becas y porcentajes", "1 (Registro Becado)", "3 (Matrícula, Porcentaje_descuento, Motivo_beca)", "Bajo", 1, 7],
        ["7. Tabla: Discapacitado", "Entidad de soporte de requerimientos y asistencia ortopédica", "1 (Registro Discapacitado)", "3 (Matrícula, Descripción_discapacidad, Aparato_ortopédico)", "Bajo", 1, 7],
    ]
    agregar_bloque_pfna("4. ARCHIVOS LÓGICOS INTERNOS (ILF) - Mapeo de las 7 Tablas de la BD",
                       ["Tabla de Base de Datos", "Propósito Funcional en el Sistema", "Registros (RET)", "Campos de la Tabla (DET)", "Complejidad", "Total Ops", "Puntos de Función"],
                       datos_ilf)

    # 5. ARCHIVOS DE INTERFAZ EXTERNA (EIF)
    datos_eif = [
        ["Validación Ciudadana (RENAPO)", "Servicio externo de validación de identidad y estado civil", "1 (Registro Ciudadano)", "4 (CURP, Nombre, Estado, Estatus)", "Bajo", 1, 5],
        ["Catálogo Homologado (SEJ / DGB)", "Servicio externo de validación de claves y programas académicos", "1 (Plan Curricular)", "3 (Clave_Oficial, Nombre, Validez)", "Bajo", 1, 5],
    ]
    agregar_bloque_pfna("5. ARCHIVOS DE INTERFAZ EXTERNA (EIF) - Integraciones Externas",
                       ["Interfaz / Catálogo Externo", "Propósito de Interoperabilidad", "Registros (RET)", "Campos Consultados (DET)", "Complejidad", "Total Ops", "Puntos de Función"],
                       datos_eif)

    # TOTAL GENERAL PFNA (Con Estilo Verde y Bordes Dobles)
    ws_pfna.merge_cells(start_row=row_p, start_column=1, end_row=row_p, end_column=6)
    tot_lbl = ws_pfna.cell(row=row_p, column=1, value="TOTAL GENERAL DE PUNTOS DE FUNCIÓN NO AJUSTADOS (PFNA):")
    tot_lbl.font = font_total_green
    tot_lbl.fill = fill_total_green
    tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
    tot_lbl.border = border_total_green

    for col in range(2, 7):
        ws_pfna.cell(row=row_p, column=col).border = border_total_green
        ws_pfna.cell(row=row_p, column=col).fill = fill_total_green

    sub_sum_formula = "+".join([f"G{r}" for r in subtotal_rows_pfna])
    tot_val = ws_pfna.cell(row=row_p, column=7)
    tot_val.value = f"={sub_sum_formula}"
    tot_val.font = font_total_green
    tot_val.fill = fill_total_green
    tot_val.alignment = Alignment(horizontal="center", vertical="center")
    tot_val.border = border_total_green
    ws_pfna.row_dimensions[row_p].height = 28
    row_total_pfna = row_p

    # Anchos de columna en Hoja 2
    col_w_p = {1: 34, 2: 44, 3: 20, 4: 28, 5: 14, 6: 12, 7: 18}
    for c_idx, width in col_w_p.items():
        ws_pfna.column_dimensions[get_column_letter(c_idx)].width = width

    # =========================================================================
    # HOJA 3: PUNTOS DE FUNCIÓN AJUSTADOS (PFA) - 14 FACTORES DE ALBRECHT
    # =========================================================================
    ws_pfa = wb.create_sheet(title="PF Ajustados")
    ws_pfa.views.sheetView[0].showGridLines = True

    # Encabezado Principal
    ws_pfa.merge_cells("A1:E1")
    t3 = ws_pfa["A1"]
    t3.value = "CÁLCULO DE PUNTOS DE FUNCIÓN AJUSTADOS (PFA) - MODELO DE ALBRECHT"
    t3.font = font_title
    t3.fill = fill_title
    t3.alignment = Alignment(horizontal="center", vertical="center")
    ws_pfa.row_dimensions[1].height = 36

    ws_pfa.merge_cells("A2:E2")
    s3 = ws_pfa["A2"]
    s3.value = "Evaluación de los 14 Factores de Ajuste (GSC) | Fórmula: PFA = PFNA * [ 0.65 + 0.01 * SUM(Fi) ]"
    s3.font = font_subtitle
    s3.alignment = Alignment(horizontal="center", vertical="center")
    ws_pfa.row_dimensions[2].height = 20

    row_a = 4

    # Cabeceras de los 14 Factores
    headers_gsc = ["N°", "Factor General del Sistema (Albrecht / GSC)", "Nivel de Influencia (0 - 5)", "Puntaje (Fi)", "Justificación Técnica en el Sistema Académico"]
    for c_i, h in enumerate(headers_gsc, 1):
        cell_h = ws_pfa.cell(row=row_a, column=c_i, value=h)
        cell_h.font = font_tbl_header
        cell_h.fill = fill_tbl_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = border_cell
    ws_pfa.row_dimensions[row_a].height = 24
    row_a += 1

    factores_albrecht = [
        ["F1", "¿Se requieren comunicaciones de datos?", "4 - Significativo", 4, "Arquitectura web con endpoints REST y sincronización HTTPS hacia la base de datos."],
        ["F2", "¿Existen funciones de procesamiento distribuido?", "3 - Moderado", 3, "Validación y procesamiento dividido entre servidor backend y lógica interactiva de interfaz."],
        ["F3", "¿Es crítico el rendimiento y performance?", "4 - Significativo", 4, "Tiempos de respuesta estrictos en cobros de caja, altas de inscripción y consultas de sábanas."],
        ["F4", "¿Será ejecutado en un entorno fuertemente utilizado?", "3 - Moderado", 3, "Operación continua con alta concurrencia en servidor local y respaldo cloud."],
        ["F5", "¿El volumen de transacciones es elevado?", "4 - Significativo", 4, "Picos intensivos durante aperturas de periodos de inscripción y pagos de colegiaturas."],
        ["F6", "¿Requiere entrada de datos interactiva en línea?", "5 - Esencial", 5, "El 100% de las operaciones de alumnos, becas, discapacidades y cobros se realiza en línea."],
        ["F7", "¿Requiere transacciones sobre múltiples pantallas?", "4 - Significativo", 4, "Múltiples vistas y pestañas operativas (kardex, materias cursadas, recibos de pago)."],
        ["F8", "¿Se actualizan los archivos maestros de forma interactiva?", "4 - Significativo", 4, "Las tablas principales (Alumno, Pago, Cursa) se actualizan en tiempo real con integridad."],
        ["F9", "¿Son complejas las entradas, salidas, archivos o peticiones?", "3 - Moderado", 3, "Validación cruzada de estados, porcentaje de becas y deducciones arancelarias."],
        ["F10", "¿Es complejo el procesamiento interno?", "3 - Moderado", 3, "Cálculo de saldos, liquidación de becas y balance financiero por alumno."],
        ["F11", "¿Se ha diseñado el código para ser reutilizable?", "4 - Significativo", 4, "Módulos de arquitectura desacoplada, controladores reutilizables y librerías estandarizadas."],
        ["F12", "¿Están incluidas la conversión y la instalación en el diseño?", "3 - Moderado", 3, "Scripts SQL DDL/DML automatizados y rutinas de migración de datos."],
        ["F13", "¿Se ha diseñado para soportar múltiples instalaciones?", "4 - Significativo", 4, "Soporta acceso multi-campus e interoperabilidad entre distintas sedes académicas."],
        ["F14", "¿Se ha diseñado la aplicación para facilitar cambios y uso?", "4 - Significativo", 4, "Diseño intuitivo, cuadrícula clara y facilidad para incorporar nuevas materias o aranceles."]
    ]

    start_gsc_r = row_a
    for idx, r_data in enumerate(factores_albrecht):
        is_even = (idx % 2 == 1)
        for c_i, val in enumerate(r_data, 1):
            cell = ws_pfa.cell(row=row_a, column=c_i, value=val)
            cell.font = font_data
            cell.border = border_cell
            if is_even:
                cell.fill = fill_zebra
            
            if c_i == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_i in [2, 5]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_pfa.row_dimensions[row_a].height = 20
        row_a += 1
    end_gsc_r = row_a - 1

    # Fila de Suma de Factores (SUM(Fi))
    ws_pfa.merge_cells(start_row=row_a, start_column=1, end_row=row_a, end_column=3)
    lbl_sum_fi = ws_pfa.cell(row=row_a, column=1, value="SUMATORIA TOTAL DE FACTORES GENERALES (SUM(Fi)):")
    lbl_sum_fi.font = font_subtotal
    lbl_sum_fi.fill = fill_subtotal
    lbl_sum_fi.alignment = Alignment(horizontal="right", vertical="center")
    lbl_sum_fi.border = border_subtotal

    for col in range(2, 4):
        ws_pfa.cell(row=row_a, column=col).border = border_subtotal
        ws_pfa.cell(row=row_a, column=col).fill = fill_subtotal

    val_sum_fi = ws_pfa.cell(row=row_a, column=4)
    val_sum_fi.value = f"=SUM(D{start_gsc_r}:D{end_gsc_r})"
    val_sum_fi.font = font_subtotal
    val_sum_fi.fill = fill_subtotal
    val_sum_fi.alignment = Alignment(horizontal="center", vertical="center")
    val_sum_fi.border = border_subtotal

    lbl_rango = ws_pfa.cell(row=row_a, column=5, value="Escala de complejidad: 0 a 70 puntos máximos")
    lbl_rango.font = Font(name="Calibri", size=9, italic=True, color="595959")
    lbl_rango.fill = fill_subtotal
    lbl_rango.alignment = Alignment(horizontal="center", vertical="center")
    lbl_rango.border = border_subtotal
    ws_pfa.row_dimensions[row_a].height = 22
    row_sum_fi = row_a
    row_a += 3

    # =========================================================================
    # CUADRO RESUMEN OFICIAL DE CÁLCULO DE PFA (Con totales verdes y fórmulas)
    # =========================================================================
    ws_pfa.merge_cells(start_row=row_a, start_column=1, end_row=row_a, end_column=5)
    t_res = ws_pfa.cell(row=row_a, column=1, value="CUADRO RESUMEN DE CÁLCULO DE PUNTOS DE FUNCIÓN AJUSTADOS (PFA)")
    t_res.font = font_section
    t_res.fill = fill_section
    t_res.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_pfa.row_dimensions[row_a].height = 24
    row_a += 1

    headers_summary = ["Parámetro / Métrica", "Valor Calculado", "Fórmula Aplicada", "Descripción e Interpretación Técnica"]
    ws_pfa.cell(row=row_a, column=1, value="Parámetro / Métrica").font = font_tbl_header
    ws_pfa.cell(row=row_a, column=1).fill = fill_tbl_header
    ws_pfa.cell(row=row_a, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws_pfa.cell(row=row_a, column=1).border = border_cell

    ws_pfa.cell(row=row_a, column=2, value="Valor").font = font_tbl_header
    ws_pfa.cell(row=row_a, column=2).fill = fill_tbl_header
    ws_pfa.cell(row=row_a, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws_pfa.cell(row=row_a, column=2).border = border_cell

    ws_pfa.cell(row=row_a, column=3, value="Fórmula de Cálculo").font = font_tbl_header
    ws_pfa.cell(row=row_a, column=3).fill = fill_tbl_header
    ws_pfa.cell(row=row_a, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws_pfa.cell(row=row_a, column=3).border = border_cell

    ws_pfa.merge_cells(start_row=row_a, start_column=4, end_row=row_a, end_column=5)
    c_desc_h = ws_pfa.cell(row=row_a, column=4, value="Descripción e Interpretación Técnica")
    c_desc_h.font = font_tbl_header
    c_desc_h.fill = fill_tbl_header
    c_desc_h.alignment = Alignment(horizontal="center", vertical="center")
    c_desc_h.border = border_cell
    ws_pfa.cell(row=row_a, column=5).border = border_cell
    ws_pfa.row_dimensions[row_a].height = 24
    row_a += 1

    # Fila 1: PFNA
    ws_pfa.cell(row=row_a, column=1, value="Puntos de Función No Ajustados (PFNA)").font = font_data_bold
    ws_pfa.cell(row=row_a, column=1).border = border_cell
    
    c_pfna = ws_pfa.cell(row=row_a, column=2, value=f"='PF No Ajustados'!G{row_total_pfna}")
    c_pfna.font = font_subtotal
    c_pfna.border = border_cell
    c_pfna.alignment = Alignment(horizontal="center", vertical="center")

    c_pfna_fml = ws_pfa.cell(row=row_a, column=3, value="SUM(EI) + SUM(EO) + SUM(EQ) + SUM(ILF) + SUM(EIF)")
    c_pfna_fml.font = font_data
    c_pfna_fml.border = border_cell
    c_pfna_fml.alignment = Alignment(horizontal="center", vertical="center")

    ws_pfa.merge_cells(start_row=row_a, start_column=4, end_row=row_a, end_column=5)
    c_pfna_d = ws_pfa.cell(row=row_a, column=4, value="Tamaño funcional intrínseco del software basado en las 7 tablas de la BD.")
    c_pfna_d.font = font_data
    c_pfna_d.border = border_cell
    ws_pfa.cell(row=row_a, column=5).border = border_cell
    ws_pfa.row_dimensions[row_a].height = 22
    row_pfna_ref = row_a
    row_a += 1

    # Fila 2: Factor de Ajuste (CAF / VAF)
    ws_pfa.cell(row=row_a, column=1, value="Factor de Ajuste de Complejidad (CAF / VAF)").font = font_data_bold
    ws_pfa.cell(row=row_a, column=1).border = border_cell
    
    c_caf = ws_pfa.cell(row=row_a, column=2, value=f"=0.65 + 0.01 * D{row_sum_fi}")
    c_caf.font = font_subtotal
    c_caf.border = border_cell
    c_caf.alignment = Alignment(horizontal="center", vertical="center")
    c_caf.number_format = "0.00"

    c_caf_fml = ws_pfa.cell(row=row_a, column=3, value=f"=0.65 + 0.01 * SUM(Fi)")
    c_caf_fml.font = font_data
    c_caf_fml.border = border_cell
    c_caf_fml.alignment = Alignment(horizontal="center", vertical="center")

    ws_pfa.merge_cells(start_row=row_a, start_column=4, end_row=row_a, end_column=5)
    c_caf_d = ws_pfa.cell(row=row_a, column=4, value="Coeficiente que pondera la complejidad operativa del entorno (rango 0.65 a 1.35).")
    c_caf_d.font = font_data
    c_caf_d.border = border_cell
    ws_pfa.cell(row=row_a, column=5).border = border_cell
    ws_pfa.row_dimensions[row_a].height = 22
    row_caf_ref = row_a
    row_a += 1

    # Fila 3: TOTAL PFA (VERDE ELEGANTE)
    ws_pfa.cell(row=row_a, column=1, value="TOTAL PUNTOS DE FUNCIÓN AJUSTADOS (PFA)").font = font_total_green
    ws_pfa.cell(row=row_a, column=1).fill = fill_total_green
    ws_pfa.cell(row=row_a, column=1).border = border_total_green
    
    c_pfa = ws_pfa.cell(row=row_a, column=2, value=f"=B{row_pfna_ref} * B{row_caf_ref}")
    c_pfa.font = font_total_green
    c_pfa.fill = fill_total_green
    c_pfa.border = border_total_green
    c_pfa.alignment = Alignment(horizontal="center", vertical="center")
    c_pfa.number_format = "#,##0.00"

    c_pfa_fml = ws_pfa.cell(row=row_a, column=3, value="=PFNA * [ 0.65 + 0.01 * SUMA(Fi) ]")
    c_pfa_fml.font = font_total_green
    c_pfa_fml.fill = fill_total_green
    c_pfa_fml.border = border_total_green
    c_pfa_fml.alignment = Alignment(horizontal="center", vertical="center")

    ws_pfa.merge_cells(start_row=row_a, start_column=4, end_row=row_a, end_column=5)
    c_pfa_d = ws_pfa.cell(row=row_a, column=4, value="Tamaño de software ajustado final utilizado para estimar esfuerzo (horas/hombre) y costo.")
    c_pfa_d.font = font_total_green
    c_pfa_d.fill = fill_total_green
    c_pfa_d.border = border_total_green
    ws_pfa.cell(row=row_a, column=5).border = border_total_green
    ws_pfa.cell(row=row_a, column=5).fill = fill_total_green
    ws_pfa.row_dimensions[row_a].height = 28

    # Anchos de columna en Hoja 3
    col_w_a = {1: 14, 2: 36, 3: 26, 4: 20, 5: 64}
    for c_idx, width in col_w_a.items():
        ws_pfa.column_dimensions[get_column_letter(c_idx)].width = width

    # =========================================================================
    # HOJA 4: RESUMEN EJECUTIVO (Consolidado y Métricas)
    # =========================================================================
    ws_res = wb.create_sheet(title="Resumen Ejecutivo")
    ws_res.views.sheetView[0].showGridLines = True

    # Encabezado Principal
    ws_res.merge_cells("A1:F1")
    t4 = ws_res["A1"]
    t4.value = "RESUMEN EJECUTIVO DE MÉTRICAS Y ESTIMACIÓN DE SOFTWARE"
    t4.font = font_title
    t4.fill = fill_title
    t4.alignment = Alignment(horizontal="center", vertical="center")
    ws_res.row_dimensions[1].height = 36

    ws_res.merge_cells("A2:F2")
    s4 = ws_res["A2"]
    s4.value = "Consolidado integral de Puntos de Función No Ajustados (PFNA) y Ajustados (PFA) | Modelo 7 Tablas"
    s4.font = font_subtitle
    s4.alignment = Alignment(horizontal="center", vertical="center")
    ws_res.row_dimensions[2].height = 20

    row_e = 4

    # Tarjetas KPI
    kpis = [
        ("PF NO AJUSTADOS", f"='PF No Ajustados'!G{row_total_pfna}", "Total de puntos funcionales puros"),
        ("FACTOR DE AJUSTE (CAF)", f"='PF Ajustados'!B{row_caf_ref}", "Ponderación de complejidad técnica"),
        ("PF AJUSTADOS (PFA)", f"='PF Ajustados'!B{row_caf_ref + 1}", "Tamaño de entrega ajustado")
    ]

    for k_idx, (kpi_t, kpi_fml, kpi_desc) in enumerate(kpis):
        c_start = (k_idx * 2) + 1
        c_end = c_start + 1
        
        ws_res.merge_cells(start_row=row_e, start_column=c_start, end_row=row_e, end_column=c_end)
        cell_lbl = ws_res.cell(row=row_e, column=c_start, value=kpi_t)
        cell_lbl.font = font_kpi_lbl
        cell_lbl.fill = fill_kpi
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.border = border_cell
        ws_res.cell(row=row_e, column=c_end).border = border_cell

        ws_res.merge_cells(start_row=row_e+1, start_column=c_start, end_row=row_e+1, end_column=c_end)
        cell_val = ws_res.cell(row=row_e+1, column=c_start, value=kpi_fml)
        cell_val.font = font_kpi_val
        cell_val.fill = fill_kpi
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.border = border_cell
        cell_val.number_format = "#,##0.00"
        ws_res.cell(row=row_e+1, column=c_end).border = border_cell

        ws_res.merge_cells(start_row=row_e+2, start_column=c_start, end_row=row_e+2, end_column=c_end)
        cell_d = ws_res.cell(row=row_e+2, column=c_start, value=kpi_desc)
        cell_d.font = Font(name="Calibri", size=8.5, italic=True, color="595959")
        cell_d.fill = fill_kpi
        cell_d.alignment = Alignment(horizontal="center", vertical="center")
        cell_d.border = border_cell
        ws_res.cell(row=row_e+2, column=c_end).border = border_cell

    ws_res.row_dimensions[row_e].height = 18
    ws_res.row_dimensions[row_e+1].height = 32
    ws_res.row_dimensions[row_e+2].height = 18
    row_e += 4

    # Tabla Resumen de Componentes Funcionales
    ws_res.merge_cells(start_row=row_e, start_column=1, end_row=row_e, end_column=6)
    t_comp = ws_res.cell(row=row_e, column=1, value="CONSOLIDADO POR TIPO DE COMPONENTE FUNCIONAL")
    t_comp.font = font_section
    t_comp.fill = fill_section
    t_comp.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_res.row_dimensions[row_e].height = 24
    row_e += 1

    headers_comp = ["Categoría / Componente", "Acrónimo", "Total de Transacciones / Archivos", "Subtotal Puntos", "Participación %", "Nivel de Complejidad Predominante"]
    for c_i, h in enumerate(headers_comp, 1):
        cell_h = ws_res.cell(row=row_e, column=c_i, value=h)
        cell_h.font = font_tbl_header
        cell_h.fill = fill_tbl_header
        cell_h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_h.border = border_cell
    ws_res.row_dimensions[row_e].height = 24
    row_e += 1

    components_summary = [
        ["Entradas Externas (Altas, Bajas, Modificaciones)", "EI", 21, f"='PF No Ajustados'!G{subtotal_rows_pfna[0]}", f"=D{row_e}/$D${row_e+5}", "Bajo / Medio"],
        ["Salidas Externas (Reportes con Cálculos)", "EO", 4, f"='PF No Ajustados'!G{subtotal_rows_pfna[1]}", f"=D{row_e+1}/$D${row_e+5}", "Medio / Alto"],
        ["Peticiones / Consultas (Búsquedas Directas)", "EQ", 7, f"='PF No Ajustados'!G{subtotal_rows_pfna[2]}", f"=D{row_e+2}/$D${row_e+5}", "Bajo / Medio"],
        ["Archivos Lógicos Internos (7 Tablas BD)", "ILF", 7, f"='PF No Ajustados'!G{subtotal_rows_pfna[3]}", f"=D{row_e+3}/$D${row_e+5}", "Bajo (7 RETs)"],
        ["Archivos de Interfaz Externa (Validaciones)", "EIF", 2, f"='PF No Ajustados'!G{subtotal_rows_pfna[4]}", f"=D{row_e+4}/$D${row_e+5}", "Bajo (2 RETs)"],
    ]

    start_comp_r = row_e
    for idx, r_data in enumerate(components_summary):
        is_even = (idx % 2 == 1)
        for c_i, val in enumerate(r_data, 1):
            cell = ws_res.cell(row=row_e, column=c_i, value=val)
            cell.font = font_data
            cell.border = border_cell
            if is_even:
                cell.fill = fill_zebra
            
            if c_i == 1:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_i == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_data_bold
            elif c_i == 5:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "0.00%"
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_res.row_dimensions[row_e].height = 20
        row_e += 1
    end_comp_r = row_e - 1

    # Total Consolidado en Hoja 4 (Verde)
    ws_res.cell(row=row_e, column=1, value="TOTAL GENERAL").font = font_total_green
    ws_res.cell(row=row_e, column=1).fill = fill_total_green
    ws_res.cell(row=row_e, column=1).border = border_total_green

    ws_res.cell(row=row_e, column=2, value="PFNA").font = font_total_green
    ws_res.cell(row=row_e, column=2).fill = fill_total_green
    ws_res.cell(row=row_e, column=2).border = border_total_green
    ws_res.cell(row=row_e, column=2).alignment = Alignment(horizontal="center", vertical="center")

    ws_res.cell(row=row_e, column=3, value=f"=SUM(C{start_comp_r}:C{end_comp_r})").font = font_total_green
    ws_res.cell(row=row_e, column=3).fill = fill_total_green
    ws_res.cell(row=row_e, column=3).border = border_total_green
    ws_res.cell(row=row_e, column=3).alignment = Alignment(horizontal="center", vertical="center")

    tot_pts = ws_res.cell(row=row_e, column=4, value=f"=SUM(D{start_comp_r}:D{end_comp_r})")
    tot_pts.font = font_total_green
    tot_pts.fill = fill_total_green
    tot_pts.border = border_total_green
    tot_pts.alignment = Alignment(horizontal="center", vertical="center")

    tot_pct = ws_res.cell(row=row_e, column=5, value=f"=SUM(E{start_comp_r}:E{end_comp_r})")
    tot_pct.font = font_total_green
    tot_pct.fill = fill_total_green
    tot_pct.border = border_total_green
    tot_pct.alignment = Alignment(horizontal="center", vertical="center")
    tot_pct.number_format = "0.00%"

    ws_res.cell(row=row_e, column=6, value="100% Cobertura").font = font_total_green
    ws_res.cell(row=row_e, column=6).fill = fill_total_green
    ws_res.cell(row=row_e, column=6).border = border_total_green
    ws_res.cell(row=row_e, column=6).alignment = Alignment(horizontal="center", vertical="center")
    ws_res.row_dimensions[row_e].height = 26

    # Anchos de columna en Hoja 4
    col_w_r = {1: 42, 2: 14, 3: 28, 4: 18, 5: 18, 6: 28}
    for c_idx, width in col_w_r.items():
        ws_res.column_dimensions[get_column_letter(c_idx)].width = width

    # Remover hoja default si existe
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    # =========================================================================
    # GUARDAR ARCHIVOS (En Downloads y en el Directorio Local)
    # =========================================================================
    downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
    filepath_downloads = os.path.join(downloads_dir, "Puntos_de_Funcion_LocalMarket.xlsx")
    filepath_local = os.path.join(os.getcwd(), "Puntos_de_Funcion_LocalMarket.xlsx")

    wb.save(filepath_downloads)
    wb.save(filepath_local)

    print(f"Archivo Excel generado con exito en: {filepath_downloads}")
    print(f"Copia de respaldo guardada en: {filepath_local}")

if __name__ == "__main__":
    generar_excel_puntos_de_funcion()
